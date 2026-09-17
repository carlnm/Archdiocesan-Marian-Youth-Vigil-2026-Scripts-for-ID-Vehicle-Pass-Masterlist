import os
import re
import pandas as pd
from PIL import Image as PILImage

# Styling & Excel generation via openpyxl
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

# ReportLab for Master and Label PDFs
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


# ============================================================
# FILE PATHS
# ============================================================

INPUT_FILE = "Parishes/Processed_Data_Cleaned.xlsx" 
CYM_INPUT_FILE = "CYM/Marian_Youth_Vigil_2026_Cleaned_School_CYM.xlsx"

OUTPUT_FOLDER = "Parishes/Marian Youth Vigil 2026 - Delegations"

DELEGATION_HEADS_FILE = os.path.join(OUTPUT_FOLDER, "Delegation Heads.xlsx")
PARISH_TRACKER_FILE = os.path.join(OUTPUT_FOLDER, "Unified_Parish_CYM_Delegation_Tracker.xlsx")
MASTER_PDF_PATH = os.path.join(OUTPUT_FOLDER, "Masterlist_All_Parishes.pdf")
LABEL_PDF_PATH = os.path.join(OUTPUT_FOLDER, "Participant_Labels_All.pdf")


# ============================================================
# PAGE SETTINGS — MASTER PDF
# ============================================================

PAGE_WIDTH = 8.5 * inch
PAGE_HEIGHT = 13.0 * inch
PAGE_SIZE = (PAGE_WIDTH, PAGE_HEIGHT)


# ============================================================
# COLORS
# ============================================================

COLOR_YELLOW_ACCENT = colors.HexColor("#FDB813")
COLOR_PRIMARY_BLUE = colors.HexColor("#0D6EFD")
COLOR_DARK_NAVY = colors.HexColor("#102A43")
COLOR_BG_LIGHT = colors.HexColor("#FFFBEB")
COLOR_BORDER = colors.HexColor("#CBD5E0")


# ============================================================
# BANNER
# ============================================================

def get_banner_info():
    possible_names = ["banner.png", "banner.jpg", "banner.jpeg", "banner"]
    for name in possible_names:
        if os.path.exists(name):
            try:
                with PILImage.open(name) as img:
                    orig_w, orig_h = img.size
                    calc_width = PAGE_WIDTH
                    calc_height = (orig_h / orig_w) * calc_width
                    return name, calc_width, calc_height
            except Exception:
                return name, PAGE_WIDTH, 1.25 * inch
    return None, 0, 0


BANNER_FILE, BANNER_WIDTH, BANNER_HEIGHT = get_banner_info()


# ============================================================
# FONTS
# ============================================================

FONT_REGULAR = "Helvetica"
FONT_BOLD = "Helvetica-Bold"


def register_garet_font():
    global FONT_REGULAR, FONT_BOLD
    garet_files = {
        "Garet-Regular": ["Garet-Book.ttf", "Garet-Regular.ttf", "Garet.ttf"],
        "Garet-Bold": ["Garet-Heavy.ttf", "Garet-Bold.ttf", "Garet-ExtraBold.ttf"]
    }

    reg_found = None
    for filename in garet_files["Garet-Regular"]:
        if os.path.exists(filename):
            pdfmetrics.registerFont(TTFont("Garet-Regular", filename))
            reg_found = "Garet-Regular"
            break

    bold_found = None
    for filename in garet_files["Garet-Bold"]:
        if os.path.exists(filename):
            pdfmetrics.registerFont(TTFont("Garet-Bold", filename))
            bold_found = "Garet-Bold"
            break

    if reg_found and bold_found:
        FONT_REGULAR = reg_found
        FONT_BOLD = bold_found
        print("Garet font registered successfully.")
    elif reg_found:
        FONT_REGULAR = reg_found
        FONT_BOLD = reg_found
        print("Garet font registered (Regular).")
    else:
        print("Garet font not found locally; falling back to Helvetica.")


register_garet_font()


# ============================================================
# NUMBERED CANVAS
# ============================================================

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        if BANNER_FILE:
            self.drawImage(
                BANNER_FILE,
                x=0,
                y=PAGE_HEIGHT - BANNER_HEIGHT,
                width=BANNER_WIDTH,
                height=BANNER_HEIGHT,
                preserveAspectRatio=True,
                mask="auto"
            )

        self.setFont(FONT_REGULAR, 8)
        self.setFillColor(colors.HexColor("#718096"))
        self.drawString(36, 20, "Archdiocesan Marian Youth Vigil 2026 — Official Delegation List")
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(PAGE_WIDTH - 36, 20, page_str)
        self.restoreState()


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def clean_name(name):
    name = str(name).strip()
    name = re.sub(r'[<>:"/\\|?*]', "_", name)
    return name.rstrip(" .")


def find_tshirt_column(df):
    possible_columns = [
        "T-Shirt Size", "Tshirt Size", "T-Shirt", "Tshirt",
        "Shirt Size", "T Shirt Size", "T shirt size", "T-shirt size"
    ]
    for col in possible_columns:
        if col in df.columns:
            return col

    for col in df.columns:
        norm = str(col).lower().replace("-", "").replace("_", "").replace(" ", "")
        if "tshirt" in norm or "shirtsize" in norm:
            return col
    return None


def find_nickname_column(df):
    possible_columns = [
        "Nickname", "Nick Name", "Nick_Name", "Alias",
        "Call Sign", "Preferred Name"
    ]
    for col in possible_columns:
        if col in df.columns:
            return col

    for col in df.columns:
        norm = str(col).lower().replace("-", "").replace("_", "").replace(" ", "")
        if "nick" in norm or "alias" in norm:
            return col
    return None


def find_parish_or_school_column(df):
    """
    Finds Parish / Church or School / CYM column reliably,
    ignoring case, spaces, and punctuation.
    """
    exact_targets = [
        "school / cym", "school/cym", "school / cy", "school/cy",
        "school", "school name", "parish / church", "parish/church",
        "parish", "church", "organization", "school / organization"
    ]
    
    col_map = {str(col).strip().lower(): col for col in df.columns}
    for target in exact_targets:
        if target in col_map:
            return col_map[target]

    for col in df.columns:
        norm = re.sub(r'[^a-z0-9]', '', str(col).lower())
        if "school" in norm or "cym" in norm or "parish" in norm or "church" in norm:
            return col

    return None


def find_delegation_head(delegation_heads_df, vicariate_name, parish_name):
    if delegation_heads_df.empty:
        return None

    target_vic = str(vicariate_name).strip().casefold()
    target_par = str(parish_name).strip().casefold()
    matched_names = []

    for _, row in delegation_heads_df.iterrows():
        row_vic = str(row.get("Vicariate", "")).strip().casefold()
        if row_vic != target_vic:
            continue

        raw_parish = str(row.get("Parish / Church", "")).strip().casefold()
        row_parishes = [p.strip() for p in raw_parish.split(";") if p.strip()]

        if target_par in row_parishes:
            name = str(row.get("Full Name", "")).strip()
            if name and not pd.isna(name):
                matched_names.append(name)

    if matched_names:
        return " / ".join(matched_names)

    return None


# ============================================================
# PARISH STORY FOR MASTER PDF
# ============================================================

def build_parish_story(parish_name, vicariate_name, records_df, nickname_col, tshirt_col, delegation_head_name):
    story_items = []
    content_width = PAGE_WIDTH - 72
    styles = getSampleStyleSheet()

    parish_style = ParagraphStyle(
        "ParishHeading",
        parent=styles["Heading1"],
        fontName=FONT_BOLD,
        fontSize=14.5,
        leading=17.5,
        textColor=COLOR_DARK_NAVY,
        spaceAfter=2
    )
    meta_left_style = ParagraphStyle(
        "MetaLeft",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=9.5,
        leading=12,
        textColor=colors.HexColor("#4A5568")
    )
    meta_right_style = ParagraphStyle(
        "MetaRight",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=9.5,
        leading=12,
        textColor=COLOR_PRIMARY_BLUE,
        alignment=TA_RIGHT
    )
    header_cell_style = ParagraphStyle(
        "HeaderCell",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=8.5,
        leading=11,
        textColor=COLOR_DARK_NAVY
    )
    header_cell_center = ParagraphStyle(
        "HeaderCellCenter",
        parent=header_cell_style,
        alignment=TA_CENTER
    )
    body_cell_style = ParagraphStyle(
        "BodyCell",
        parent=styles["Normal"],
        fontName=FONT_REGULAR,
        fontSize=8.5,
        leading=11,
        textColor=COLOR_DARK_NAVY
    )
    body_cell_center = ParagraphStyle(
        "BodyCellCenter",
        parent=body_cell_style,
        alignment=TA_CENTER
    )
    body_cell_shirt = ParagraphStyle(
        "BodyCellShirt",
        parent=body_cell_style,
        fontName=FONT_BOLD,
        textColor=COLOR_PRIMARY_BLUE,
        alignment=TA_CENTER
    )

    story_items.append(Paragraph(str(parish_name).upper(), parish_style))

    meta_table_data = [[
        Paragraph(f"VICARIATE: {vicariate_name}", meta_left_style),
        Paragraph(f"TOTAL DELEGATES: {len(records_df)}", meta_right_style)
    ]]

    meta_table = Table(meta_table_data, colWidths=[350, content_width - 350])
    meta_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LINEBELOW", (0, 0), (-1, -1), 2.5, COLOR_YELLOW_ACCENT),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0)
    ]))

    story_items.append(meta_table)
    story_items.append(Spacer(1, 8))

    col_widths = [35, 255, 150, 100]
    table_data = [[
        Paragraph("#", header_cell_center),
        Paragraph("FULL NAME", header_cell_style),
        Paragraph("NICKNAME", header_cell_style),
        Paragraph("T-SHIRT", header_cell_center)
    ]]

    for idx, (_, row) in enumerate(records_df.iterrows(), start=1):
        fullname = str(row.get("Full Name", "")).strip() or "-"
        nickname = str(row.get(nickname_col, "")).strip() or "-" if nickname_col else "-"
        tshirt = str(row.get(tshirt_col, "")).strip() or "-" if tshirt_col else "-"

        table_data.append([
            Paragraph(str(idx), body_cell_center),
            Paragraph(fullname, body_cell_style),
            Paragraph(nickname, body_cell_style),
            Paragraph(tshirt, body_cell_shirt)
        ])

    participant_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    t_style = [
        ("BACKGROUND", (0, 0), (-1, 0), COLOR_YELLOW_ACCENT),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, COLOR_BORDER),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6)
    ]

    for i in range(1, len(table_data)):
        bg = COLOR_BG_LIGHT if i % 2 == 0 else colors.white
        t_style.append(("BACKGROUND", (0, i), (-1, i), bg))

    participant_table.setStyle(TableStyle(t_style))
    story_items.append(participant_table)
    story_items.append(Spacer(1, 20))

    signature_title_style = ParagraphStyle(
        "SignatureTitle",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=9.5,
        leading=12,
        textColor=COLOR_DARK_NAVY
    )
    signature_name_style = ParagraphStyle(
        "SignatureName",
        parent=styles["Normal"],
        fontName=FONT_BOLD,
        fontSize=10,
        leading=13,
        textColor=COLOR_DARK_NAVY,
        alignment=TA_CENTER
    )
    signature_label_style = ParagraphStyle(
        "SignatureLabel",
        parent=styles["Normal"],
        fontName=FONT_REGULAR,
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#718096"),
        alignment=TA_CENTER
    )

    display_name = delegation_head_name.upper() if delegation_head_name else "________________________________"

    signature_table = Table([
        [Paragraph("DELEGATION HEAD", signature_title_style)],
        [Spacer(1, 24)],
        [Paragraph(display_name, signature_name_style)],
        [Paragraph("Signature over Printed Name", signature_label_style)],
        [Spacer(1, 8)],
        [Paragraph("Date: ______________________________", signature_label_style)]
    ], colWidths=[250], hAlign="LEFT")

    signature_table.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2)
    ]))

    story_items.append(signature_table)
    return story_items


# ============================================================
# UNIFIED INTERACTIVE TRACKER (PARISH + CYM)
# ============================================================

def build_unified_interactive_excel(all_participants_df, all_delegation_heads_df, output_path):
    print()
    print("==============================================")
    print("BUILDING UNIFIED INTERACTIVE TRACKER (PARISH + CYM)...")
    print("==============================================")

    # 1. Count Participants
    p_counts = (
        all_participants_df.groupby(["Category", "Vicariate", "Parish / Church"])
        .size()
        .reset_index(name="Participants")
    )

    # 2. Count Delegation Heads
    head_rows = []
    if all_delegation_heads_df is not None and not all_delegation_heads_df.empty:
        for _, row in all_delegation_heads_df.iterrows():
            cat = str(row.get("Category", "Parish")).strip()
            vic = str(row.get("Vicariate", "Not Specified")).strip()
            raw_parish = str(row.get("Parish / Church", "Not Specified")).strip()
            parish_list = [p.strip() for p in raw_parish.split(";") if p.strip()]
            if not parish_list:
                parish_list = ["Not Specified"]
            for p in parish_list:
                head_rows.append({"Category": cat, "Vicariate": vic, "Parish / Church": p})

    heads_expanded = pd.DataFrame(head_rows)
    if not heads_expanded.empty:
        h_counts = (
            heads_expanded.groupby(["Category", "Vicariate", "Parish / Church"])
            .size()
            .reset_index(name="Delegation Heads")
        )
    else:
        h_counts = pd.DataFrame(columns=["Category", "Vicariate", "Parish / Church", "Delegation Heads"])

    # 3. Merge both counts
    merged = pd.merge(p_counts, h_counts, on=["Category", "Vicariate", "Parish / Church"], how="outer")
    merged["Participants"] = merged["Participants"].fillna(0).astype(int)
    merged["Delegation Heads"] = merged["Delegation Heads"].fillna(0).astype(int)

    # Sort: Category -> Vicariate -> Parish / Church
    merged = merged.sort_values(
        by=["Category", "Vicariate", "Parish / Church"],
        key=lambda col: col.astype(str).str.casefold()
    ).reset_index(drop=True)

    num_rows = len(merged)
    first_data_row = 7
    last_data_row = 6 + num_rows
    total_row = last_data_row + 1

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Delegation Tracker"

    # Freeze top 6 rows
    ws.freeze_panes = "A7"

    thin_border = Border(
        left=Side(style="thin", color="CBD5E0"),
        right=Side(style="thin", color="CBD5E0"),
        top=Side(style="thin", color="CBD5E0"),
        bottom=Side(style="thin", color="CBD5E0")
    )
    title_fill = PatternFill(start_color="102A43", end_color="102A43", fill_type="solid")
    hdr_fill = PatternFill(start_color="FDB813", end_color="FDB813", fill_type="solid")
    card_hdr_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
    total_fill = PatternFill(start_color="102A43", end_color="102A43", fill_type="solid")
    banner_base_fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")

    # Row 1: Title Header
    ws.merge_cells("A1:G1")
    ws["A1"] = "ARCHDIOCESAN MARIAN YOUTH VIGIL 2026 — UNIFIED PARISH & CYM DELEGATION TRACKER"
    ws["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    # Row 2: Dynamic Alert Formula Banner (Monitors Column G: Total Pax)
    ws.merge_cells("A2:G2")
    ws["A2"] = (
        f'=IF(SUMIF(A{first_data_row}:A{last_data_row}, TRUE, G{first_data_row}:G{last_data_row})>=400, '
        f'"⚠️ NOTIFICATION: TARGET REACHED! " & TEXT(SUMIF(A{first_data_row}:A{last_data_row}, TRUE, G{first_data_row}:G{last_data_row}), "#,##0") & " / 400 TOTAL DELEGATES SELECTED", '
        f'"STATUS: " & TEXT(SUMIF(A{first_data_row}:A{last_data_row}, TRUE, G{first_data_row}:G{last_data_row}), "#,##0") & " / 400 PAX SELECTED (" & '
        f'TEXT(MAX(0, 400 - SUMIF(A{first_data_row}:A{last_data_row}, TRUE, G{first_data_row}:G{last_data_row})), "#,##0") & " SLOTS REMAINING)")'
    )
    ws["A2"].font = Font(name="Segoe UI", size=11, bold=True, color="2D3748")
    ws["A2"].fill = banner_base_fill
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 26

    # Rows 3 & 4: Quick Summary Cards
    ws["A3"] = "CHECKED BODIES"
    ws["A4"] = f'=COUNTIF(A{first_data_row}:A{last_data_row}, TRUE)'

    ws.merge_cells("B3:D3")
    ws["B3"] = "DELEGATION OVERVIEW"
    ws.merge_cells("B4:D4")
    ws["B4"] = f"Total {num_rows} Parishes & CYM/School Units Loaded"

    ws["E3"] = "SELECTED PARTICIPANTS"
    ws["E4"] = f'=SUMIF(A{first_data_row}:A{last_data_row}, TRUE, E{first_data_row}:E{last_data_row})'

    ws["F3"] = "SELECTED HEADS"
    ws["F4"] = f'=SUMIF(A{first_data_row}:A{last_data_row}, TRUE, F{first_data_row}:F{last_data_row})'

    ws["G3"] = "CHECKED TOTAL PAX"
    ws["G4"] = f'=SUMIF(A{first_data_row}:A{last_data_row}, TRUE, G{first_data_row}:G{last_data_row})'

    for cell_ref in ["A3", "E3", "F3", "G3", "B3"]:
        ws[cell_ref].font = Font(name="Segoe UI", size=8.5, bold=True, color="4A5568")
        ws[cell_ref].fill = card_hdr_fill
        ws[cell_ref].alignment = Alignment(horizontal="center", vertical="center")
        ws[cell_ref].border = thin_border

    for cell_ref in ["A4", "E4", "F4"]:
        ws[cell_ref].font = Font(name="Segoe UI", size=13, bold=True, color="0D6EFD")
        ws[cell_ref].alignment = Alignment(horizontal="center", vertical="center")
        ws[cell_ref].border = thin_border

    ws["B4"].font = Font(name="Segoe UI", size=9.5, italic=True, color="718096")
    ws["B4"].alignment = Alignment(horizontal="center", vertical="center")
    ws["B4"].border = thin_border

    ws["G4"].font = Font(name="Segoe UI", size=14, bold=True, color="102A43")
    ws["G4"].fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
    ws["G4"].alignment = Alignment(horizontal="center", vertical="center")
    ws["G4"].border = Border(
        left=Side(style="medium", color="FDB813"),
        right=Side(style="medium", color="FDB813"),
        top=Side(style="medium", color="FDB813"),
        bottom=Side(style="medium", color="FDB813")
    )

    ws.row_dimensions[3].height = 16
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 8

    # Row 6: Table Headers
    headers = [
        "CHECK (TRUE/FALSE)",
        "CATEGORY",
        "VICARIATE",
        "PARISH / SCHOOL / ORGANIZATION",
        "PARTICIPANTS",
        "DELEGATION HEADS",
        "TOTAL PAX"
    ]
    for c_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=6, column=c_idx, value=h_text)
        cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="102A43")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[6].height = 22

    # Dropdown Validation for TRUE / FALSE
    dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"A{first_data_row}:A{last_data_row}")

    # Data Rows
    for idx, row in merged.iterrows():
        r = first_data_row + idx
        ws.cell(row=r, column=1, value=False).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=2, value=str(row["Category"])).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=3, value=str(row["Vicariate"])).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=4, value=str(row["Parish / Church"])).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r, column=5, value=int(row["Participants"])).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=r, column=6, value=int(row["Delegation Heads"])).alignment = Alignment(horizontal="center", vertical="center")

        # Row Total Pax = Col E + Col F
        tot_cell = ws.cell(row=r, column=7, value=f"=E{r}+F{r}")
        tot_cell.font = Font(name="Segoe UI", size=9.5, bold=True, color="0D6EFD")
        tot_cell.alignment = Alignment(horizontal="center", vertical="center")

        for col in range(1, 8):
            ws.cell(row=r, column=col).border = thin_border
        ws.row_dimensions[r].height = 20

    # Grand Total Row
    ws.cell(row=total_row, column=1, value="ALL").alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=total_row, column=2, value="TOTAL").alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=total_row, column=3, value="-").alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=total_row, column=4, value=f"All {num_rows} Units Combined").alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(row=total_row, column=5, value=f"=SUM(E{first_data_row}:E{last_data_row})").alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=total_row, column=6, value=f"=SUM(F{first_data_row}:F{last_data_row})").alignment = Alignment(horizontal="center", vertical="center")

    tot_pax = ws.cell(row=total_row, column=7, value=f"=SUM(G{first_data_row}:G{last_data_row})")
    tot_pax.alignment = Alignment(horizontal="center", vertical="center")
    tot_pax.font = Font(name="Segoe UI", size=11, bold=True, color="102A43")

    for col in range(1, 8):
        c = ws.cell(row=total_row, column=col)
        c.fill = total_fill if col <= 4 else hdr_fill
        c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF" if col <= 4 else "102A43")
        c.border = thin_border
    ws.row_dimensions[total_row].height = 24

    # Conditional Formatting: Banner turns red when checked total >= 400
    red_fill = PatternFill(start_color="FED7D7", end_color="FED7D7", fill_type="solid")
    red_font = Font(name="Segoe UI", size=11, bold=True, color="9B2C2C")
    ws.conditional_formatting.add(
        "A2:G2",
        FormulaRule(
            formula=[f'SUMIF(A{first_data_row}:A{last_data_row}, TRUE, G{first_data_row}:G{last_data_row})>=400'],
            fill=red_fill,
            font=red_font
        )
    )

    # Column Widths
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 44
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 16

    wb.save(output_path)
    print(f"Interactive Parish + CYM Tracker Excel created: {output_path}")


# ============================================================
# PARTICIPANT LABEL SETTINGS
# ============================================================

LABEL_PAGE_WIDTH = 8.5 * inch
LABEL_PAGE_HEIGHT = 11.0 * inch
LABEL_PAGE_SIZE = (LABEL_PAGE_WIDTH, LABEL_PAGE_HEIGHT)

LABEL_COLUMNS = 3
LABEL_ROWS = 8
LABELS_PER_PAGE = LABEL_COLUMNS * LABEL_ROWS

LABEL_MARGIN_X = 0.20 * inch
LABEL_MARGIN_Y = 0.20 * inch
LABEL_GAP_X = 0.06 * inch
LABEL_GAP_Y = 0.06 * inch

LABEL_WIDTH = (
    LABEL_PAGE_WIDTH - (2 * LABEL_MARGIN_X) - ((LABEL_COLUMNS - 1) * LABEL_GAP_X)
) / LABEL_COLUMNS

LABEL_HEIGHT = (
    LABEL_PAGE_HEIGHT - (2 * LABEL_MARGIN_Y) - ((LABEL_ROWS - 1) * LABEL_GAP_Y)
) / LABEL_ROWS


# ============================================================
# PARTICIPANT LABEL PDF (INCLUDES DELEGATION HEADS)
# ============================================================

def build_participant_labels_pdf(participants_df, delegation_heads_df, output_path):
    print()
    print("==============================================")
    print("GENERATING COMBINED LABEL PDF (PARTICIPANTS + HEADS)...")
    print("==============================================")

    p_df = participants_df.copy()
    p_df["Full Name"] = p_df["Full Name"].fillna("").astype(str).str.strip()
    p_df["Parish / Church"] = p_df["Parish / Church"].fillna("Not Specified").astype(str).str.strip()
    p_df["Vicariate"] = p_df["Vicariate"].fillna("Not Specified").astype(str).str.strip()
    p_df["Is_Head"] = False
    p_df = p_df[p_df["Full Name"] != ""].copy()

    head_records = []
    if delegation_heads_df is not None and not delegation_heads_df.empty:
        for _, row in delegation_heads_df.iterrows():
            name = str(row.get("Full Name", "")).strip()
            vic = str(row.get("Vicariate", "Not Specified")).strip()
            raw_parishes = str(row.get("Parish / Church", "Not Specified")).strip()

            if not name:
                continue

            parish_list = [p.strip() for p in raw_parishes.split(";") if p.strip()]
            if not parish_list:
                parish_list = ["Not Specified"]

            for p in parish_list:
                head_records.append({
                    "Full Name": name,
                    "Parish / Church": p,
                    "Vicariate": vic,
                    "Is_Head": True
                })

    heads_expanded_df = pd.DataFrame(head_records)
    combined_df = pd.concat([heads_expanded_df, p_df], ignore_index=True)

    combined_df = combined_df.sort_values(
        by=["Vicariate", "Parish / Church", "Is_Head", "Full Name"],
        ascending=[True, True, False, True],
        key=lambda col: col if col.name == "Is_Head" else col.astype(str).str.casefold()
    ).reset_index(drop=True)

    c = canvas.Canvas(output_path, pagesize=LABEL_PAGE_SIZE)

    NAME_FONT = FONT_BOLD
    PARISH_FONT = FONT_REGULAR
    VICARIATE_FONT = FONT_REGULAR
    ROLE_FONT = FONT_BOLD

    def fit_text(text, font_name, font_size, max_width, minimum_size=5):
        text = str(text).strip()
        current_size = font_size

        while current_size >= minimum_size:
            if pdfmetrics.stringWidth(text, font_name, current_size) <= max_width:
                return text, current_size
            current_size -= 0.25

        current_size = minimum_size
        ellipsis = "..."
        shortened_text = text

        while len(shortened_text) > 4:
            shortened = shortened_text[:-4] + ellipsis
            if pdfmetrics.stringWidth(shortened, font_name, current_size) <= max_width:
                return shortened, current_size
            shortened_text = shortened_text[:-1]

        return shortened_text, current_size

    def draw_centered_text(text, center_x, y, font_name, font_size, color=COLOR_DARK_NAVY):
        c.saveState()
        c.setFont(font_name, font_size)
        c.setFillColor(color)
        c.drawCentredString(center_x, y, text)
        c.restoreState()

    for index, row in combined_df.iterrows():
        position = index % LABELS_PER_PAGE
        column = position % LABEL_COLUMNS
        row_number = position // LABEL_COLUMNS

        x = LABEL_MARGIN_X + column * (LABEL_WIDTH + LABEL_GAP_X)
        y = LABEL_PAGE_HEIGHT - LABEL_MARGIN_Y - ((row_number + 1) * LABEL_HEIGHT) - (row_number * LABEL_GAP_Y)

        c.setStrokeColor(colors.HexColor("#A0AEC0"))
        c.setLineWidth(0.5)
        c.rect(x, y, LABEL_WIDTH, LABEL_HEIGHT, stroke=1, fill=0)

        center_x = x + LABEL_WIDTH / 2
        padding = 0.08 * inch
        max_text_width = LABEL_WIDTH - (2 * padding)

        fullname = row["Full Name"]
        parish = row["Parish / Church"]
        vicariate = row["Vicariate"]
        is_head = row.get("Is_Head", False)

        if is_head:
            display_name, name_size = fit_text(fullname, NAME_FONT, 8.5, max_text_width, minimum_size=6.5)
            display_parish, parish_size = fit_text(parish, PARISH_FONT, 6.0, max_text_width, minimum_size=5.0)
            display_vicariate, vicariate_size = fit_text(vicariate, VICARIATE_FONT, 6.0, max_text_width, minimum_size=5.0)

            name_y = y + LABEL_HEIGHT - 0.28 * inch
            role_y = y + LABEL_HEIGHT - 0.44 * inch
            parish_y = y + LABEL_HEIGHT - 0.68 * inch
            vicariate_y = y + LABEL_HEIGHT - 0.88 * inch

            draw_centered_text(display_name, center_x, name_y, NAME_FONT, name_size)
            draw_centered_text("[ DELEGATION HEAD ]", center_x, role_y, ROLE_FONT, 6.0, color=COLOR_PRIMARY_BLUE)
            draw_centered_text(display_parish, center_x, parish_y, PARISH_FONT, parish_size)
            draw_centered_text(display_vicariate, center_x, vicariate_y, VICARIATE_FONT, vicariate_size)

        else:
            display_name, name_size = fit_text(fullname, NAME_FONT, 9.5, max_text_width, minimum_size=7)
            display_parish, parish_size = fit_text(parish, PARISH_FONT, 6.5, max_text_width, minimum_size=5.5)
            display_vicariate, vicariate_size = fit_text(vicariate, VICARIATE_FONT, 6.5, max_text_width, minimum_size=5.5)

            name_y = y + LABEL_HEIGHT - 0.36 * inch
            parish_y = y + LABEL_HEIGHT - 0.63 * inch
            vicariate_y = y + LABEL_HEIGHT - 0.84 * inch

            draw_centered_text(display_name, center_x, name_y, NAME_FONT, name_size)
            draw_centered_text(display_parish, center_x, parish_y, PARISH_FONT, parish_size)
            draw_centered_text(display_vicariate, center_x, vicariate_y, VICARIATE_FONT, vicariate_size)

        if position == (LABELS_PER_PAGE - 1):
            c.showPage()

    if len(combined_df) % LABELS_PER_PAGE != 0:
        c.showPage()

    c.save()

    total_pages = (len(combined_df) + LABELS_PER_PAGE - 1) // LABELS_PER_PAGE
    print("Combined Label PDF created:")
    print(f"  Participants: {len(p_df)}")
    print(f"  Delegation Head Labels: {len(heads_expanded_df)}")
    print(f"  Total Labels: {len(combined_df)}")
    print(f"  Pages: {total_pages}")
    print(f"  File: {output_path}")


# ============================================================
# MAIN PROGRAM
# ============================================================

def main():
    print()
    print("==============================================")
    print("MARIAN YOUTH VIGIL 2026")
    print("PARISH & CYM DELEGATION PROCESSOR")
    print("==============================================")

    os.makedirs(OUTPUT_FOLDER, exist_ok=True)

    # 1. Load Parish Excel
    print(f"Reading Parish Excel: {INPUT_FILE}...")
    df_parish = pd.read_excel(INPUT_FILE)
    df_parish.columns = [str(col).strip() for col in df_parish.columns]
    df_parish["Category"] = "Parish"

    parish_col = find_parish_or_school_column(df_parish)
    if parish_col and parish_col != "Parish / Church":
        df_parish["Parish / Church"] = df_parish[parish_col]

    # 2. Load CYM Excel
    df_cym = pd.DataFrame()
    if os.path.exists(CYM_INPUT_FILE):
        print(f"Reading CYM Excel: {CYM_INPUT_FILE}...")
        df_cym = pd.read_excel(CYM_INPUT_FILE)
        df_cym.columns = [str(col).strip() for col in df_cym.columns]
        df_cym["Category"] = "CYM / School"

        cym_col = find_parish_or_school_column(df_cym)
        print(f"--> Detected CYM School/Org column: '{cym_col}'")

        if cym_col:
            df_cym["Parish / Church"] = df_cym[cym_col]
        else:
            print("⚠️ WARNING: Could not find School/CYM column in CYM file!")

        print(f"Loaded {len(df_cym)} CYM rows.")
    else:
        print(f"WARNING: CYM file '{CYM_INPUT_FILE}' not found.")

    # Combine both datasets
    dfs_to_combine = [df_parish]
    if not df_cym.empty:
        dfs_to_combine.append(df_cym)

    df_combined = pd.concat(dfs_to_combine, ignore_index=True)

    # Clean strings and ensure valid text
    for col in ["Full Name", "Parish / Church", "Vicariate"]:
        if col in df_combined.columns:
            df_combined[col] = (
                df_combined[col]
                .fillna("Not Specified")
                .astype(str)
                .str.strip()
                .replace("", "Not Specified")
            )

    # Determine Type column (Participant vs Delegation Head)
    type_col = None
    for col in df_combined.columns:
        norm = str(col).lower().replace("_", "").replace("-", "").replace(" ", "")
        if norm == "type":
            type_col = col
            break

    if type_col:
        all_participants = df_combined[
            df_combined[type_col].astype(str).str.strip().str.casefold() == "participant"
        ].copy()
        all_delegation_heads = df_combined[
            df_combined[type_col].astype(str).str.strip().str.casefold().isin([
                "delegation head", "delegationhead", "head"
            ])
        ].copy()
    else:
        print("WARNING: No Type column found. Treating all rows as participants.")
        all_participants = df_combined.copy()
        all_delegation_heads = pd.DataFrame()

    all_participants = all_participants[all_participants["Full Name"] != ""].copy()

    nickname_col = find_nickname_column(df_combined)
    tshirt_col = find_tshirt_column(df_combined)

    if not all_delegation_heads.empty:
        all_delegation_heads.to_excel(DELEGATION_HEADS_FILE, index=False)
        print(f"Delegation Heads Excel created: {DELEGATION_HEADS_FILE}")

    # Generate Unified Interactive Tracker Excel
    build_unified_interactive_excel(
        all_participants_df=all_participants,
        all_delegation_heads_df=all_delegation_heads,
        output_path=PARISH_TRACKER_FILE
    )

    # Parish individual folder Excels (only from original parish list)
    parish_participants = all_participants[all_participants["Category"] == "Parish"].copy()
    parish_groups = parish_participants.groupby(["Vicariate", "Parish / Church"], dropna=False)
    for ((vicariate, parish), p_df) in parish_groups:
        safe_vicariate = clean_name(vicariate)
        safe_parish = clean_name(parish)
        vic_folder = os.path.join(OUTPUT_FOLDER, safe_vicariate)
        os.makedirs(vic_folder, exist_ok=True)
        out_file = os.path.join(vic_folder, f"{safe_parish}.xlsx")
        p_df.to_excel(out_file, index=False)

    # Master PDF (Parishes)
    master_pdf_story = []
    sorted_parishes = parish_participants.sort_values(
        by=["Vicariate", "Parish / Church", "Full Name"],
        key=lambda col: col.astype(str).str.casefold()
    )
    grouped = sorted_parishes.groupby(["Vicariate", "Parish / Church"], sort=False)

    first_parish = True
    for ((vicariate, parish), records_df) in grouped:
        if not first_parish:
            master_pdf_story.append(PageBreak())
        first_parish = False

        delegation_head_name = find_delegation_head(all_delegation_heads, vicariate, parish)
        master_pdf_story.extend(
            build_parish_story(
                parish_name=parish,
                vicariate_name=vicariate,
                records_df=records_df,
                nickname_col=nickname_col,
                tshirt_col=tshirt_col,
                delegation_head_name=delegation_head_name
            )
        )

    if master_pdf_story:
        doc = SimpleDocTemplate(
            MASTER_PDF_PATH,
            pagesize=PAGE_SIZE,
            rightMargin=36,
            leftMargin=36,
            topMargin=(BANNER_HEIGHT + 18),
            bottomMargin=36
        )
        doc.build(master_pdf_story, canvasmaker=NumberedCanvas)
        print(f"Master PDF created: {MASTER_PDF_PATH}")

    # T-Shirt Summary
    if tshirt_col:
        tshirt_summary = (
            all_participants[tshirt_col]
            .fillna("Not Specified")
            .astype(str)
            .str.strip()
            .replace("", "Not Specified")
            .value_counts()
            .rename_axis("T-Shirt Size")
            .reset_index(name="Quantity")
        )
        tshirt_summary_path = os.path.join(OUTPUT_FOLDER, "T-Shirt Summary.xlsx")
        tshirt_summary.to_excel(tshirt_summary_path, index=False)

    # Business Insights
    insights_path = os.path.join(OUTPUT_FOLDER, "Business Insights.xlsx")
    with pd.ExcelWriter(insights_path, engine="openpyxl") as writer:
        vicariate_summary = (
            all_participants.groupby("Vicariate")
            .size()
            .reset_index(name="Participants")
            .sort_values("Participants", ascending=False)
        )
        vicariate_summary.to_excel(writer, sheet_name="Vicariate Summary", index=False)

        parish_summary = (
            all_participants.groupby(["Category", "Vicariate", "Parish / Church"])
            .size()
            .reset_index(name="Participants")
            .sort_values("Participants", ascending=False)
        )
        parish_summary.to_excel(writer, sheet_name="Unit Summary", index=False)

        overall_summary = pd.DataFrame({
            "Metric": ["Total Participants", "Total Vicariates", "Total Units (Parish + CYM)"],
            "Value": [len(all_participants), all_participants["Vicariate"].nunique(), all_participants["Parish / Church"].nunique()]
        })
        overall_summary.to_excel(writer, sheet_name="Overall Summary", index=False)

    # Labels PDF
    build_participant_labels_pdf(
        participants_df=all_participants,
        delegation_heads_df=all_delegation_heads,
        output_path=LABEL_PDF_PATH
    )

    print()
    print("==============================================")
    print("PROCESS COMPLETED SUCCESSFULLY")
    print("==============================================")


if __name__ == "__main__":
    main()